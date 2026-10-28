"""Build paper/research_paper.docx from paper/main.tex (pandoc), then highlight every [xyz] in yellow.

    python scripts/build_docx.py
Re-run after editing main.tex or replacing figures. Equations become editable Word equations,
citations are formatted from references.bib, figures use the PNG versions.
"""

from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"


def preprocess(tex: str) -> str:
    tex = re.sub(r"^\\newcommand\{\\xyz\}.*$", "", tex, flags=re.M)
    tex = re.sub(r"^\\newcommand\{\\feat\}.*$", "", tex, flags=re.M)
    tex = re.sub(r"\\xyz(\[[^\]]*\])?(\{\})?", "[xyz]", tex)  # keyed placeholders, e.g. \xyz[h1.bfree.delta]{}
    tex = re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]*)\.pdf\}", r"\\includegraphics\1{\2.png}", tex)
    tex = re.sub(r"\\feat\{([^}]*)\}", r"\\texttt{\1}", tex)
    tex = re.sub(r"\\multirow\{[^}]*\}\{[^}]*\}\{([^}]*)\}", r"\1", tex)
    tex = tex.replace("\\textperiodcentered ", "\u00b7").replace("\\textperiodcentered", "\u00b7")
    tex = tex.replace("\\\\ A Controlled", ": A Controlled")
    return tex


def highlight(docx_path: Path) -> int:
    doc = Document(docx_path)
    count = 0

    def paragraphs():
        for p in doc.paragraphs:
            yield p
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    for p in paragraphs():
        if "[xyz]" not in p.text:
            continue
        for run in list(p.runs):
            if "[xyz]" not in run.text:
                continue
            parts = re.split(r"(\[xyz\])", run.text)
            run.text = parts[0]
            anchor = run
            for part in parts[1:]:
                if not part:
                    continue
                new = anchor._element.__class__  # noqa: F841
                r = p.add_run(part)
                r.bold, r.italic = run.bold, run.italic
                r.font.name, r.font.size = run.font.name, run.font.size
                anchor._element.addnext(r._element)
                anchor = r
                if part == "[xyz]":
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    count += 1
    for style_name in ("Normal", "Body Text", "First Paragraph"):
        try:
            st = doc.styles[style_name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(11)
        except KeyError:
            pass
    doc.save(docx_path)
    return count


def main() -> int:
    tex = preprocess((PAPER / "main.tex").read_text())
    out = PAPER / "research_paper.docx"
    with tempfile.NamedTemporaryFile("w", suffix=".tex", dir=PAPER, delete=False) as f:
        f.write(tex)
        tmp = Path(f.name)
    try:
        cmd = ["pandoc", tmp.name, "-o", str(out), "--citeproc", "--bibliography=references.bib",
               "--resource-path=.:figures", "--number-sections", "-M", "reference-section-title=References",
               "-M", "link-citations=true"]
        subprocess.run(cmd, cwd=PAPER, check=True)
    finally:
        tmp.unlink()
    n = highlight(out)
    print(f"Wrote {out} with {n} yellow [xyz] placeholders")
    return 0


if __name__ == "__main__":
    sys.exit(main())
