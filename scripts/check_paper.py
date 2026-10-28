"""Check the manuscript before submission.

1. Hyphenated words in the PROSE of main.tex and supplementary.tex (the style guide asks for none).
   Official names (B-Free, CO-SPY, SSD-1B, FLUX.1-dev, ...), identifiers in \\texttt, citation keys,
   labels, URLs and math are ignored.
2. Count of remaining [xyz] placeholders (LaTeX \\xyz and DOCX "[xyz]").
3. Figures still marked as placeholders (paper/figures/*.placeholder).
4. Citation keys used in the .tex files that are missing from references.bib.

    python scripts/check_paper.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"

ALLOWED = {
    "B-Free", "CO-SPY", "Co-Spy-Bench", "SSD-1B", "tiny-sd", "small-sd", "FLUX.1-dev", "FLUX.1-schnell", "ViT-L/14",
    "MoGe-2", "DRCT-2M", "AI-GenBench", "Qwen2.5-VL-7B-Instruct", "Qwen2.5-VL-7B", "LAION-400M", "RAISE-1k", "DINOv2-L",
    "Clarke-Pearson", "SD-1", "sd-v1", "D\\&B", "scikit-learn", "ViT-L",
}


def strip_latex(text: str) -> str:
    text = re.sub(r"(?<!\\)%.*", "", text)                      # comments
    text = re.sub(r"\$[^$]*\$", " ", text)                       # inline math
    text = re.sub(r"\\begin\{(equation|align)\*?\}.*?\\end\{\1\*?\}", " ", text, flags=re.S)
    text = re.sub(r"\\(cite[pt]?|citet|citep|label|ref|url|href|includegraphics|texttt|feat|bibliography|bibliographystyle|graphicspath|usepackage|documentclass|newcommand|renewcommand)\*?(\[[^\]]*\])?\{[^}]*\}", " ", text)
    text = re.sub(r"\\begin\{[^}]*\}(\{[^}]*\})?|\\end\{[^}]*\}", " ", text)
    return text


def hyphen_words(text: str) -> list[str]:
    words = [w.strip(".") for w in re.findall(r"[A-Za-z0-9.\\&/]+(?:-[A-Za-z0-9.]+)+", text)]
    return [w for w in words if w not in ALLOWED and not re.fullmatch(r"[\d.\-]+", w)]


def main() -> int:
    problems = 0
    bib = (PAPER / "references.bib").read_text()
    bib_keys = set(re.findall(r"@\w+\{([^,]+),", bib))
    for name in ("main.tex", "supplementary.tex"):
        p = PAPER / name
        if not p.exists():
            continue
        raw = p.read_text()
        body = raw.split("\\begin{document}", 1)[-1]
        hw = hyphen_words(strip_latex(body))
        n_xyz = len(re.findall(r"\\xyz", body))
        used = set(k.strip() for grp in re.findall(r"\\cite[pt]?\{([^}]*)\}", body) for k in grp.split(","))
        missing = sorted(used - bib_keys)
        print(f"{name}: {n_xyz} [xyz] placeholders left")
        if hw:
            problems += 1
            print(f"  hyphenated words in prose ({len(hw)}): {sorted(set(hw))}")
        else:
            print("  no hyphenated words in prose")
        if missing:
            problems += 1
            print(f"  citation keys missing from references.bib: {missing}")
    docx = PAPER / "research_paper.docx"
    if docx.exists():
        try:
            from docx import Document

            d = Document(docx)
            text = "\n".join(par.text for par in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " ".join(c.text for c in row.cells)
            print(f"research_paper.docx: {text.count('[xyz]')} [xyz] placeholders left")
        except ImportError:
            print("python-docx not installed; skipped DOCX check")
    ph = sorted(p.stem for p in (PAPER / "figures").glob("*.placeholder"))
    print(f"figures still placeholders ({len(ph)}): {ph}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
